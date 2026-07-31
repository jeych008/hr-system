#!/usr/bin/env python3
import hashlib
import json
import re
import secrets
import string
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


PROJECTS = [
    ("p_jiangsu_inbound", "江苏呼入", "hr_why"),
    ("p_jiangsu_outbound", "江苏呼出", "hr_clf"),
    ("p_suzhou_downgrade", "苏州降档", "hr_clf"),
    ("p_nanjing_downgrade", "南京降档", "hr_clf"),
    ("p_shanghai_songjiang", "上海-松江", "hr_wxy"),
    ("p_shanghai_luoyang", "上海-洛阳", "hr_pxj"),
    ("p_meiyu_guangxi", "美毓-广西", "hr_xhj"),
    ("p_shuozhou", "朔州职场", "hr_wrx"),
    ("p_gaoyou", "高邮职场", "hr_qyx"),
    ("p_huaian_pool", "淮安资源池", "hr_yln"),
    ("p_nanjing_post", "南京邮政", "hr_ld"),
    ("p_wuxi_post", "无锡邮政", "hr_ld"),
    ("p_shaanxi_weinan", "陕西渭南", "hr_hb"),
]

HR_NAMES = {
    "hr_why": "王海园", "hr_clf": "陈龙凤", "hr_wxy": "吴心雨",
    "hr_pxj": "潘兴姣", "hr_xhj": "邢惠娟", "hr_wrx": "王瑞雪",
    "hr_qyx": "钱雨欣", "hr_yln": "姚林娜", "hr_ld": "李丹", "hr_hb": "黄冰",
}

SOURCE_PROJECTS = {
    "淮安清河呼入职场": "p_jiangsu_inbound",
    "淮安清河外呼职场": "p_jiangsu_outbound",
    "南京六合职场": "p_nanjing_downgrade",
    "上海松江职场": "p_shanghai_songjiang",
    "洛阳西工职场": "p_shanghai_luoyang",
    "美毓临汾职场": "p_meiyu_guangxi",
    "美毓通信临汾职场": "p_meiyu_guangxi",
    "德汇朔州职场": "p_shuozhou",
    "高邮职场": "p_gaoyou",
    "淮安资源池职场": "p_huaian_pool",
    "无锡新业务事业部": "p_wuxi_post",
    "陕西渭南职场": "p_shaanxi_weinan",
}

STATUS_MAP = {
    "已签到": "ARRIVED",
    "已录用": "PASSED",
    "面试不通过": "FAILED",
    "岗前培训中": "TRAINING",
    "已离职": "LEFT",
}


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def new_id(prefix):
    return f"{prefix}_{uuid4().hex[:16]}"


def valid_id_card(value):
    value = text(value).upper()
    if not re.fullmatch(r"[0-9]{17}[0-9X]", value):
        return False
    try:
        datetime.strptime(value[6:14], "%Y%m%d")
    except ValueError:
        return False
    weights = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
    checks = "10X98765432"
    return checks[sum(int(value[index]) * weights[index] for index in range(17)) % 11] == value[-1]


def birth_from_id(value):
    value = text(value).upper()
    return f"{value[6:10]}-{value[10:12]}-{value[12:14]}"


def gender_from_id(value):
    return "男" if int(text(value)[16]) % 2 else "女"


def valid_phone(value):
    return bool(re.fullmatch(r"1[3-9][0-9]{9}", text(value)))


def excel_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = text(value).replace("/", "-").replace(".", "-")
    match = re.fullmatch(r"(20[0-9]{2}|19[0-9]{2})-([0-9]{1,2})-([0-9]{1,2})", raw)
    if not match:
        return ""
    try:
        return date(int(match[1]), int(match[2]), int(match[3])).isoformat()
    except ValueError:
        return ""


def normalize_education(value):
    value = text(value)
    return {
        "硕士": "研究生", "中专": "中专或技校", "技校": "中专或技校",
        "初中": "初中及以下", "小学": "初中及以下",
    }.get(value, value or "未知")


def normalize_mandarin(value):
    value = text(value)
    return {
        "无等级": "未评级", "": "未评级", "一级甲等": "一甲", "一级乙等": "一乙",
        "二级甲等": "二甲", "二级乙等": "二乙", "三级甲等": "三甲", "三级乙等": "三乙",
    }.get(value, "未评级")


def normalize_channel(value):
    raw = text(value)
    compact = re.sub(r"[\s　]+", "", raw).lower()
    if "58" in compact:
        return "58同城"
    if "boss" in compact or "boos" in compact or compact in {"bss", "直聘"}:
        return "BOSS直聘"
    if any(token in raw for token in ["朋友", "介绍", "推荐"]):
        return f"亲友介绍：{raw}"
    if any(token in raw for token in ["人才市场", "招聘会"]):
        return "现场招聘会"
    if any(token in raw for token in ["招聘网", "招聘网站", "本地"]):
        return "本地招聘网"
    return f"其他：{raw or '历史数据未提供'}"


def normalize_location(value):
    raw = text(value) or "历史数据未提供"
    parts = [part for part in re.split(r"[/\\,，;；\s]+", raw) if part]
    province = parts[0] if parts else raw
    city = parts[1] if len(parts) > 1 else province
    return province, city, raw


def parse_work_experiences(value, fallback_position):
    raw = text(value)
    if not raw:
        return []
    header = re.compile(
        r"(?m)^(?P<company>[^\n|]*?)\s*\|\s*(?P<position>[^\n(，,]*?)"
        r"(?:\s*\((?P<start>[0-9]{4}-[0-9]{2}-[0-9]{2})\s*~\s*(?P<end>[0-9]{4}-[0-9]{2}-[0-9]{2})\))?"
        r"[，,]\s*工作内容[：:]\s*"
    )
    matches = list(header.finditer(raw))
    if not matches:
        return [{
            "companyName": "历史数据未提供", "position": fallback_position,
            "startDate": "", "endDate": "", "description": raw,
        }]
    experiences = []
    for index, match in enumerate(matches):
        end_offset = matches[index + 1].start() if index + 1 < len(matches) else len(raw)
        description = raw[match.end():end_offset].strip()
        start_date = match.group("start") or ""
        end_date = match.group("end") or ""
        if start_date and end_date and end_date < start_date:
            description = f"原始时间：{start_date} ~ {end_date}\n{description}".strip()
            start_date = ""
            end_date = ""
        company = match.group("company").strip()
        position = match.group("position").strip()
        experiences.append({
            "companyName": company if company and company != "-" else "历史数据未提供",
            "position": position if position and position != "-" else fallback_position,
            "startDate": start_date,
            "endDate": end_date,
            "description": description or "历史数据未提供",
        })
    return experiences


def random_password():
    alphabet = string.ascii_letters + string.digits + "!@#$%"
    while True:
        password = "".join(secrets.choice(alphabet) for _ in range(16))
        if re.search(r"[A-Z]", password) and re.search(r"[a-z]", password) and re.search(r"[0-9]", password):
            return password


def password_hash(password):
    salt = secrets.token_hex(16)
    digest = hashlib.scrypt(password.encode(), salt=salt.encode(), n=16384, r=8, p=1, dklen=32)
    return f"scrypt:{salt}:{digest.hex()}"


def style_sheet(sheet, landscape=False):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    fill = PatternFill("solid", fgColor="173B6C")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in sheet.columns:
        width = min(42, max(12, max(len(text(cell.value)) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if landscape else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.sheet_view.showGridLines = False


def write_workbook(path, title, headers, rows, landscape=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    style_sheet(sheet, landscape=landscape)
    workbook.save(path)
    path.chmod(0o600)


def candidate_from_record(record, project_id, hr_id):
    id_card = text(record["身份证号"]).upper()
    phone = text(record["手机号"])
    source_gender = text(record["性别"])
    gender = gender_from_id(id_card) if source_gender == "未知" else source_gender
    residence_province, residence_city, residence = normalize_location(record["现住地址"])
    expected_province, expected_city, expected = normalize_location(record["意向工作地点"])
    applied_position = text(record["应聘岗位"]) or "历史数据未提供"
    available_date = excel_date(record["可到岗时间"])
    graduation_time = excel_date(record["毕业时间"])
    status = STATUS_MAP[text(record["状态"])]
    arrived = status in {"ARRIVED", "PASSED", "FAILED", "TRAINING", "LEFT"}
    passed = status in {"PASSED", "TRAINING", "LEFT"}
    employment_status = status if status in {"TRAINING", "LEFT"} else "CANDIDATE"
    failed_at = "1970-01-01T00:00:00.000Z" if status == "FAILED" else None
    work_experiences = parse_work_experiences(record["工作经历"], applied_position)
    return {
        "id": new_id("cand"), "projectId": project_id, "source": "PRODUCTION_IMPORT", "createdBy": hr_id,
        "createdAt": "1970-01-01T00:00:00.000Z", "updatedAt": "1970-01-01T00:00:00.000Z",
        "name": text(record["姓名"]), "gender": gender, "birthDate": birth_from_id(id_card),
        "education": normalize_education(record["最高学历"]), "phone": phone, "idCard": id_card,
        "email": "", "recruitmentChannel": normalize_channel(record["招聘渠道"]),
        "appliedPosition": applied_position, "availableDate": available_date,
        "graduationSchool": text(record["学校名称"]) or "历史数据未提供",
        "major": text(record["专业"]) or "历史数据未提供", "mandarinLevel": normalize_mandarin(record["普通话等级"]),
        "graduationTime": graduation_time, "phoneCustomerServiceExperience": text(record["客服经验"]) in {"是", "有"},
        "hasCustomerServiceExperience": text(record["客服经验"]) in {"是", "有"},
        "currentResidenceProvince": residence_province, "currentResidenceCity": residence_city, "currentResidence": residence,
        "maritalStatus": text(record["婚姻状况"]) or "未知", "childrenStatus": "有" if text(record["是否有子女"]) == "是" else "无",
        "expectedProvince": expected_province, "expectedCity": expected_city, "expectedLocation": expected,
        "lastLeaveReason": text(record["离职原因"]) or "历史数据未提供",
        "workExperiences": work_experiences,
        "arrived": arrived, "arrivedAt": None, "arrivedBy": hr_id if arrived else None,
        "passed": passed, "passedAt": None, "passedBy": hr_id if passed else None,
        "employmentStatus": employment_status, "plannedJoinDate": available_date if status == "PASSED" else None,
        "plannedJoinSetAt": None, "plannedJoinSetBy": hr_id if status == "PASSED" else None,
        "failedAt": failed_at, "failedBy": hr_id if status == "FAILED" else None,
        "failedReason": "历史数据：面试不通过" if status == "FAILED" else None,
        "joinedAt": None, "joinedBy": hr_id if status in {"TRAINING", "LEFT"} else None,
        "trainingStartedAt": None, "trainingStartedBy": hr_id if status == "TRAINING" else None,
        "trainingEndedAt": None, "trainingEndedBy": None, "activeAt": None, "activeBy": None,
        "leftAt": None, "leftBy": hr_id if status == "LEFT" else None,
        "leftReason": "未获取" if status == "LEFT" else None, "leftFromStatus": None,
        "abandonAt": None, "abandonBy": None, "abandonReason": None, "lastReminderDate": None,
        "historicalData": True, "historicalStatusDateUnknown": status in {"TRAINING", "LEFT"},
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: prepare-production-import.py <source.xlsx> <output-directory>")
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    output.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["累计候选人信息"]
    headers = [text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    records = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True) if any(value not in (None, "") for value in values)]

    duplicate_keys = Counter()
    for record in records:
        project_id = SOURCE_PROJECTS.get(text(record["所属项目组"]))
        duplicate_keys[(project_id, text(record["手机号"]))] += 1

    project_by_id = {project_id: (name, username) for project_id, name, username in PROJECTS}
    projects = [{
        "id": project_id, "name": name, "shortCode": f"job-{hashlib.sha256(project_id.encode()).hexdigest()[:8]}",
        "enabled": True, "createdAt": datetime.now().astimezone().isoformat(),
    } for project_id, name, _ in PROJECTS]
    project_ids_by_hr = defaultdict(list)
    for project_id, _, username in PROJECTS:
        project_ids_by_hr[username].append(project_id)

    credentials = []
    users = []
    for username, display_name in HR_NAMES.items():
        password = random_password()
        user_id = f"u_{username}"
        credentials.append([display_name, username, password, "首次登录后请立即修改密码"])
        users.append({
            "id": user_id, "username": username, "displayName": display_name, "role": "HR", "enabled": True,
            "passwordHash": password_hash(password), "authVersion": 0, "projectIds": project_ids_by_hr[username],
        })

    candidates = []
    accepted_rows = []
    review = []
    status_counts = Counter()
    project_counts = Counter()
    for row_number, record in enumerate(records, start=2):
        issues = []
        source_project = text(record["所属项目组"])
        project_id = SOURCE_PROJECTS.get(source_project)
        phone = text(record["手机号"])
        id_card = text(record["身份证号"]).upper()
        if not project_id:
            issues.append("项目无法映射")
        if not valid_phone(phone):
            issues.append("手机号不是有效11位中国大陆手机号")
        if not valid_id_card(id_card):
            issues.append("身份证不是有效18位身份证号")
        if project_id and duplicate_keys[(project_id, phone)] > 1:
            issues.append("同项目手机号重复，整组暂不导入")
        if text(record["状态"]) not in STATUS_MAP:
            issues.append("状态无法映射")
        if issues:
            review.append([row_number, "；".join(issues)] + [text(record.get(header)) for header in headers])
            continue
        _, hr_username = project_by_id[project_id]
        candidate = candidate_from_record(record, project_id, f"u_{hr_username}")
        candidates.append(candidate)
        cleaned = dict(record)
        cleaned["性别"] = candidate["gender"]
        cleaned["所属项目组"] = project_by_id[project_id][0]
        cleaned["状态"] = {
            "ARRIVED": "已到面", "PASSED": "已通过", "FAILED": "未通过", "TRAINING": "培训中", "LEFT": "已离职"
        }[STATUS_MAP[text(record["状态"])]]
        accepted_rows.append([text(cleaned.get(header)) for header in headers])
        status_counts[text(record["状态"])] += 1
        project_counts[project_id] += 1

    generated_at = datetime.now().astimezone().isoformat()
    audit_logs = [{
        "id": new_id("log"), "actorId": candidate["createdBy"], "action": "PRODUCTION_IMPORT",
        "entityType": "CANDIDATE", "entityId": candidate["id"], "before": None, "after": candidate, "createdAt": generated_at,
    } for candidate in candidates]
    bundle = {
        "manifest": {
            "formatVersion": 1, "generatedAt": generated_at, "sourceFile": source.name, "sourceRows": len(records),
            "acceptedRows": len(candidates), "reviewRows": len(review), "projectCount": len(projects), "hrCount": len(users),
            "statusCounts": dict(status_counts), "projectCounts": dict(project_counts),
        },
        "projects": projects, "monthlyConfigs": [], "users": users, "candidates": candidates, "auditLogs": audit_logs,
    }
    bundle_path = output / "production-import.json"
    bundle_path.write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")
    bundle_path.chmod(0o600)
    write_workbook(output / "HR临时账号.xlsx", "HR临时账号", ["HR姓名", "用户名", "临时密码", "说明"], credentials)
    write_workbook(output / "已清洗可导入数据.xlsx", "已清洗可导入数据", headers, accepted_rows, landscape=True)
    write_workbook(output / "待修正数据.xlsx", "待修正数据", ["原始行号", "问题"] + headers, review, landscape=True)
    summary_rows = [["源数据行数", len(records)], ["可导入行数", len(candidates)], ["待修正行数", len(review)], ["项目数", len(projects)], ["HR账号数", len(users)]]
    summary_rows += [[f"状态：{key}", value] for key, value in sorted(status_counts.items())]
    write_workbook(output / "导入汇总.xlsx", "导入汇总", ["指标", "数量"], summary_rows)
    print(json.dumps(bundle["manifest"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

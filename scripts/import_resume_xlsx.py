import json
import re
import sys
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


ROOT = Path(__file__).resolve().parents[1]
DB_FILE = ROOT / "data" / "db.json"
DEFAULT_PROJECT_ID = "p_service_center"
DEFAULT_USER_ID = "u_admin"


def new_id(prefix):
    return f"{prefix}_{uuid4().hex[:16]}"


def text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    value = str(value).strip()
    if value.endswith(".0") and value[:-2].isdigit():
        return value[:-2]
    return value


def excel_date(value):
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except Exception:
            return text(value)
    return text(value)


def birth_from_id(id_card):
    value = text(id_card).upper()
    if re.match(r"^\d{17}[\dX]$", value):
        return f"{value[6:10]}-{value[10:12]}-{value[12:14]}"
    if re.match(r"^\d{15}$", value):
        return f"19{value[6:8]}-{value[8:10]}-{value[10:12]}"
    return "1990-01-01"


def split_location(value):
    parts = [p for p in re.split(r"[/\\\n\s]+", text(value)) if p and p not in {"无", "未", "暂无"}]
    if not parts:
        return "", "", ""
    province = parts[0]
    city = parts[1] if len(parts) > 1 else parts[0]
    return province, city, " ".join(parts)


def yes_no(value):
    return text(value) in {"是", "有", "Y", "yes", "true", "True", "1"}


def children(value):
    return "有" if yes_no(value) else "无"


def normalize_education(value):
    value = text(value)
    mapping = {
        "初中": "初中及以下",
        "小学": "初中及以下",
        "初中及以下": "初中及以下",
        "高中/中专": "中专或技校",
        "高中 / 中专": "中专或技校",
        "技校中专": "中专或技校",
        "职高": "中专或技校",
        "技校": "中专或技校",
        "专科": "大专",
        "高中": "高中",
        "中专": "中专或技校",
        "大专": "大专",
        "本科": "本科",
        "硕士": "研究生",
        "研究生": "研究生",
        "博士": "博士",
    }
    return mapping.get(value, "大专")


def normalize_mandarin(value):
    value = text(value)
    mapping = {
        "一级甲等": "一甲",
        "一级乙等": "一乙",
        "一级及以上": "一乙",
        "二级甲等": "二甲",
        "二级乙等": "二乙",
        "二级以上": "二乙",
        "三级甲等": "三甲",
        "三级乙等": "三乙",
        "无等级": "未评级",
        "无": "未评级",
        "未标注": "未评级",
        "-": "未评级",
    }
    if value in {"未评级", "一甲", "一乙", "二甲", "二乙", "三甲", "三乙"}:
        return value
    return mapping.get(value, "未评级")


def normalize_channel(value):
    value = text(value)
    if value in {"BOSS直聘", "58同城", "本地招聘网", "现场招聘会", "社区推荐"}:
        return value
    if "58" in value:
        return "58同城"
    if "社区" in value:
        return "社区推荐"
    if "现场" in value or "人才市场" in value or "招聘会" in value:
        return "现场招聘会"
    if "朋友" in value or "亲友" in value or "介绍" in value or "推荐" in value:
        return f"亲友介绍：{value}"
    if "BOSS" in value.upper() or "BOSS" in value or "boss" in value:
        return "BOSS直聘"
    if "招聘网站" in value or "本地" in value:
        return "本地招聘网"
    return f"其他：{value or '未标注'}"


def filled(value, fallback):
    return text(value) or fallback


def load_db():
    return json.loads(DB_FILE.read_text(encoding="utf-8"))


def save_db(db):
    DB_FILE.write_text(json.dumps(db, ensure_ascii=False, indent=2), encoding="utf-8")


def import_file(xlsx_path, project_id=DEFAULT_PROJECT_ID):
    db = load_db()
    project = next((p for p in db["projects"] if p["id"] == project_id), None)
    if not project:
        raise SystemExit(f"项目不存在: {project_id}")

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    headers = [text(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    existing_phones = {c.get("phone") for c in db.get("candidates", []) if c.get("projectId") == project_id}
    imported = 0
    skipped = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        name = text(record.get("请填写姓名（必填）"))
        phone = text(record.get("请填写手机号（必填）"))
        id_card = text(record.get("请填写身份证号（必填）")).upper()
        if not name or not phone or phone in existing_phones:
            skipped += 1
            continue

        residence_province, residence_city, residence_full = split_location(record.get("请填写现住址（必填）"))
        expected_province, expected_city, expected_full = split_location(record.get("意向工作地点（必填）"))
        residence_province = residence_province or "上海市"
        residence_city = residence_city or "上海市"
        expected_province = expected_province or "上海市"
        expected_city = expected_city or "上海市"
        applied_position = text(record.get("应聘岗位（必填）")) or "客服"
        work_description = "Excel导入工作经历"
        leave_reasons = text(record.get("请选择您离职的具体原因（可多选）：（必填）")) or "Excel导入数据未填写"
        submitted = record.get("提交时间（自动）")
        created_at = submitted.isoformat() if isinstance(submitted, datetime) else datetime.now().isoformat()

        candidate = {
            "id": new_id("cand"),
            "projectId": project_id,
            "source": "IMPORT",
            "createdBy": DEFAULT_USER_ID,
            "createdAt": created_at,
            "updatedAt": datetime.now().isoformat(),
            "name": name,
            "gender": text(record.get("请选择性别（必填）")) or "其他",
            "birthDate": birth_from_id(id_card),
            "education": normalize_education(record.get("请选择最高学历（必填）")),
            "phone": phone,
            "idCard": id_card,
            "email": f"{phone}@import.local",
            "recruitmentChannel": normalize_channel(record.get("* 获知招聘信息渠道（必填）")),
            "appliedPosition": applied_position,
            "availableDate": excel_date(record.get("可到岗时间（必填）")) or "随时到岗",
            "graduationSchool": filled(record.get("请填写毕业院校（必填）"), "未填写"),
            "major": filled(record.get("所学专业（必填）"), "未填写"),
            "mandarinLevel": normalize_mandarin(record.get("您的普通话水平等级是？（必填）")),
            "graduationTime": filled(record.get("请选择毕业时间（必填）"), "未填写"),
            "hasCustomerServiceExperience": yes_no(record.get("是否有电话客服类行业经验（必填）")),
            "phoneCustomerServiceExperience": yes_no(record.get("是否有电话客服类行业经验（必填）")),
            "currentResidenceProvince": residence_province,
            "currentResidenceCity": residence_city,
            "currentResidence": residence_full or f"{residence_province} {residence_city}",
            "maritalStatus": text(record.get("您的婚姻状况是？（必填）")) if text(record.get("您的婚姻状况是？（必填）")) in {"已婚", "未婚"} else "未婚",
            "childrenStatus": children(record.get("您是否有子女？（必填）")),
            "expectedProvince": expected_province,
            "expectedCity": expected_city,
            "expectedLocation": expected_full or f"{expected_province} {expected_city}",
            "lastLeaveReason": leave_reasons,
            "workExperiences": [
                {
                    "companyName": "Excel导入",
                    "position": applied_position,
                    "startDate": text(record.get("请选择毕业时间（必填）")) or "未知",
                    "endDate": excel_date(record.get("可到岗时间（必填）")) or "未知",
                    "description": work_description,
                }
            ],
            "arrived": False,
            "arrivedAt": None,
            "arrivedBy": None,
            "passed": False,
            "passedAt": None,
            "passedBy": None,
        }
        db["candidates"].append(candidate)
        db["auditLogs"].insert(0, {
            "id": new_id("log"),
            "actorId": DEFAULT_USER_ID,
            "action": "IMPORT_EXCEL",
            "entityType": "CANDIDATE",
            "entityId": candidate["id"],
            "before": None,
            "after": candidate,
            "createdAt": datetime.now().isoformat(),
        })
        existing_phones.add(phone)
        imported += 1

    save_db(db)
    print(json.dumps({"imported": imported, "skipped": skipped, "totalCandidates": len(db["candidates"])}, ensure_ascii=False))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit("Usage: import_resume_xlsx.py <xlsx_path> [project_id]")
    import_file(Path(sys.argv[1]), sys.argv[2] if len(sys.argv) > 2 else DEFAULT_PROJECT_ID)
